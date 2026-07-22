"""Server-side XLSX export for analytics dashboard.

Layout:
  - Resumen (period, KPIs, conclusions)
  - One sheet per indicator with native Excel chart + table + rich analysis
  - Hidden _sistema metadata
"""

from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, DoughnutChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from app.core.time import APP_TIMEZONE_NAME, now_colombia
from app.schemas.analytics_v2 import (
    ANALYTICS_SCHEMA_VERSION,
    AnalyticsModule,
    DashboardResponse,
    VisualizationType,
)
from app.services.analytics_analysis import analysis_parameter_rows, build_rich_analysis
from app.services.analytics_colors import (
    breakdown_color,
    domain_color,
    occupancy_color,
    ranking_color,
)

# Forbidden technical terms on visible sheets (regression guard).
FORBIDDEN_VISIBLE_TERMS = (
    "module_id",
    "series_id",
    "schema_version",
    "query_key",
    "trace_id",
    "payload",
    "cache_age_seconds",
    "enum",
)

_HEADER_FILL = PatternFill(start_color="1A1C1C", end_color="1A1C1C", fill_type="solid")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_TITLE_FONT = Font(bold=True, size=14)
_SECTION_FONT = Font(bold=True, size=12)
_MAX_COL_WIDTH = 48
_SHEET_TITLE_MAX = 31


def _logo_candidates() -> list[Path]:
    root = Path(__file__).resolve().parents[4]  # repo root from apps/api/app/services
    return [
        root / "apps" / "mobile" / "assets" / "branding" / "lajuana.svg",
        root / "apps" / "mobile" / "assets" / "branding" / "splash_logo.png",
        root / "apps" / "mobile" / "assets" / "branding" / "lajuana-banner.svg",
    ]


def _set_widths(ws: Any, widths: list[float]) -> None:
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = min(width, _MAX_COL_WIDTH)


def _style_header(ws: Any, row: int = 1) -> None:
    for cell in ws[row]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="center")


def _try_table(ws: Any, name: str, ref: str) -> None:
    try:
        safe = re.sub(r"[^A-Za-z0-9_]", "_", name)[:50] or "Tabla"
        table = Table(displayName=safe, ref=ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)
    except Exception:
        # Tables are enhancement; never fail export.
        pass


def _sheet_title(title: str, used: set[str]) -> str:
    base = re.sub(r"[\\/*?:\[\]]", "-", title).strip() or "Indicador"
    base = base[:_SHEET_TITLE_MAX]
    candidate = base
    n = 2
    while candidate in used:
        suffix = f" ({n})"
        candidate = (base[: _SHEET_TITLE_MAX - len(suffix)] + suffix)[:_SHEET_TITLE_MAX]
        n += 1
    used.add(candidate)
    return candidate


def _category_label(category: Any) -> str:
    raw = category.value if hasattr(category, "value") else str(category)
    return {
        "action": "Atención",
        "reservations": "Reservas",
        "money": "Dinero",
        "experiences": "Experiencias",
        "participants": "Participantes",
        "equines": "Equinos",
        "operations": "Operación",
    }.get(raw, raw)


def _apply_doughnut_colors(chart: DoughnutChart, colors: list[str]) -> None:
    if not chart.series:
        return
    series = chart.series[0]
    series.data_points = []
    for i, hex_color in enumerate(colors):
        pt = DataPoint(idx=i)
        pt.graphicalProperties.solidFill = hex_color
        series.data_points.append(pt)


def _apply_bar_colors(chart: BarChart, colors: list[str]) -> None:
    if not chart.series:
        return
    series = chart.series[0]
    series.data_points = []
    for i, hex_color in enumerate(colors):
        pt = DataPoint(idx=i)
        pt.graphicalProperties.solidFill = hex_color
        series.data_points.append(pt)


def _apply_line_color(chart: LineChart, hex_color: str) -> None:
    if not chart.series:
        return
    series = chart.series[0]
    series.graphicalProperties.line.solidFill = hex_color
    series.graphicalProperties.line.width = 25000  # EMUs (~2pt)


class AnalyticsExportService:
    def export_dashboard(
        self,
        dashboard: DashboardResponse,
        *,
        generated_by: str,
        include_logo: bool = True,
    ) -> tuple[bytes, str]:
        wb = Workbook()

        self._write_resumen(wb, dashboard, generated_by=generated_by, include_logo=include_logo)
        used_titles: set[str] = {"Resumen", "_sistema"}
        for mod in dashboard.modules:
            self._write_module_sheet(wb, mod, used_titles=used_titles)
        self._write_hidden_meta(wb, dashboard, generated_by=generated_by)

        try:
            wb.properties.title = "Reporte de analítica — La Juana"
            wb.properties.creator = generated_by
            wb.properties.description = (
                f"schema={ANALYTICS_SCHEMA_VERSION};"
                f"period={dashboard.period.start.isoformat()}..{dashboard.period.end.isoformat()}"
            )
        except Exception:
            pass

        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        data = buf.getvalue()

        stamp = now_colombia().strftime("%Y%m%d_%H%M")
        filename = (
            f"analitica_{dashboard.period.start.isoformat()}_"
            f"{dashboard.period.end.isoformat()}_{stamp}.xlsx"
        )
        return data, filename

    def _write_resumen(
        self,
        wb: Workbook,
        dashboard: DashboardResponse,
        *,
        generated_by: str,
        include_logo: bool,
    ) -> None:
        ws = wb.active
        ws.title = "Resumen"
        ws.append(["Reporte de analítica — La Juana"])
        ws["A1"].font = _TITLE_FONT
        ws.append(["Periodo", dashboard.period.label])
        ws.append(
            [
                "Desde",
                dashboard.period.start.isoformat(),
                "Hasta",
                dashboard.period.end.isoformat(),
            ]
        )
        local_now = now_colombia()
        ws.append(
            [
                "Elaborado",
                local_now.strftime("%d/%m/%Y %H:%M"),
                "Zona",
                APP_TIMEZONE_NAME.replace("_", " "),
            ]
        )
        ws.append(["Elaborado por", generated_by])
        ws.append([])
        ws.append(["Principales cifras"])
        ws["A7"].font = _SECTION_FONT
        for mod in dashboard.modules[:8]:
            if mod.primary_value is None:
                continue
            ws.append(
                [
                    mod.title,
                    mod.primary_value.formatted,
                    mod.primary_value.unit,
                    mod.insight_text or "",
                ]
            )
        ws.append([])
        ws.append(["Conclusiones"])
        for mod in dashboard.modules:
            if mod.insight_text:
                ws.append([mod.title, mod.insight_text])

        if include_logo:
            for candidate in _logo_candidates():
                if not candidate.exists():
                    continue
                if candidate.suffix.lower() == ".svg":
                    continue
                try:
                    from openpyxl.drawing.image import Image as XLImage

                    img = XLImage(str(candidate))
                    img.width = 120
                    img.height = 40
                    ws.add_image(img, "F1")
                    break
                except Exception:
                    continue

        _set_widths(ws, [36, 24, 18, 64])

    def _write_module_sheet(
        self,
        wb: Workbook,
        mod: AnalyticsModule,
        *,
        used_titles: set[str],
    ) -> None:
        title = _sheet_title(mod.title, used_titles)
        ws = wb.create_sheet(title)

        ws.append([mod.title])
        ws["A1"].font = _TITLE_FONT
        ws.append(["Categoría", _category_label(mod.category)])
        ws.append(["Periodo", mod.period.label])
        if mod.primary_value:
            ws.append(
                [
                    "Valor principal",
                    mod.primary_value.formatted,
                    mod.primary_value.unit,
                ]
            )
        if mod.description:
            ws.append(["Descripción", mod.description])
        ws.append([])

        # Key parameters (quantitative snapshot)
        params = analysis_parameter_rows(mod)
        if params:
            ws.append(["Parámetros"])
            ws.cell(row=ws.max_row, column=1).font = _SECTION_FONT
            header_row = ws.max_row + 1
            ws.append(["Parámetro", "Valor"])
            _style_header(ws, header_row)
            for key, value in params:
                ws.append([key, value])
            _try_table(
                ws,
                f"Params_{mod.id}",
                f"A{header_row}:B{ws.max_row}",
            )
            ws.append([])

        # Analysis section (conversational paragraphs)
        ws.append(["Análisis"])
        analysis_row = ws.max_row
        ws.cell(row=analysis_row, column=1).font = _SECTION_FONT
        paragraphs = list(mod.analysis) if mod.analysis else build_rich_analysis(mod)
        for line in paragraphs:
            ws.append([line])
        ws.append([])

        # Data + chart
        viz = mod.visualization
        cat_raw = mod.category.value if hasattr(mod.category, "value") else str(mod.category)

        if mod.series and viz in (
            VisualizationType.LINE,
            VisualizationType.SPARKLINE,
            VisualizationType.KPI,
        ):
            self._write_series_block(ws, mod, accent=domain_color(cat_raw))
        elif viz == VisualizationType.DONUT or (
            mod.breakdown and viz not in (VisualizationType.ACTION_LIST,)
        ):
            self._write_breakdown_block(ws, mod)
        elif mod.ranking or viz in (
            VisualizationType.RANKING,
            VisualizationType.PROGRESS,
            VisualizationType.BAR,
        ):
            self._write_ranking_block(ws, mod)
        elif mod.breakdown:
            self._write_breakdown_block(ws, mod, chart=False)
        else:
            ws.append(["Sin datos tabulares para este indicador."])

        _set_widths(ws, [36, 18, 18, 18, 18])

    def _write_series_block(
        self,
        ws: Any,
        mod: AnalyticsModule,
        *,
        accent: str,
    ) -> None:
        series = mod.series[0] if mod.series else None
        if series is None or not series.points:
            ws.append(["Sin puntos de serie."])
            return

        ws.append(["Datos de la serie"])
        ws.cell(row=ws.max_row, column=1).font = _SECTION_FONT
        header_row = ws.max_row + 1
        ws.append(["Fecha o periodo", "Valor", "Unidad"])
        _style_header(ws, header_row)
        data_start = header_row + 1
        for point in series.points:
            ws.append([point.label, point.raw, point.unit or series.unit])
        data_end = ws.max_row
        _try_table(
            ws,
            f"Serie_{mod.id}",
            f"A{header_row}:C{data_end}",
        )

        chart = LineChart()
        chart.title = series.label or mod.title
        chart.style = 10
        chart.y_axis.title = series.unit or ""
        chart.height = 10
        chart.width = 18
        data_ref = Reference(ws, min_col=2, min_row=header_row, max_row=data_end)
        cats = Reference(ws, min_col=1, min_row=data_start, max_row=data_end)
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats)
        _apply_line_color(chart, accent)
        # Place chart to the right of the table
        ws.add_chart(chart, "E" + str(header_row))

    def _write_breakdown_block(
        self,
        ws: Any,
        mod: AnalyticsModule,
        *,
        chart: bool = True,
    ) -> None:
        items = list(mod.breakdown)
        if not items:
            ws.append(["Sin desglose."])
            return

        ws.append(["Desglose"])
        ws.cell(row=ws.max_row, column=1).font = _SECTION_FONT
        header_row = ws.max_row + 1
        ws.append(["Nombre", "Cantidad", "Participación %", "Posición"])
        _style_header(ws, header_row)
        data_start = header_row + 1
        colors: list[str] = []
        for idx, item in enumerate(items):
            ws.append(
                [
                    item.label,
                    item.raw_value,
                    item.share_percentage,
                    item.rank,
                ]
            )
            colors.append(breakdown_color(mod.id, item.key, idx))
        data_end = ws.max_row
        _try_table(ws, f"Desglose_{mod.id}", f"A{header_row}:D{data_end}")

        if not chart:
            return

        pie = DoughnutChart()
        pie.title = mod.title
        labels = Reference(ws, min_col=1, min_row=data_start, max_row=data_end)
        data_ref = Reference(ws, min_col=2, min_row=header_row, max_row=data_end)
        pie.add_data(data_ref, titles_from_data=True)
        pie.set_categories(labels)
        pie.dataLabels = DataLabelList()
        pie.dataLabels.showPercent = True
        pie.dataLabels.showVal = False
        pie.dataLabels.showCatName = False
        pie.height = 10
        pie.width = 14
        _apply_doughnut_colors(pie, colors)
        ws.add_chart(pie, "F" + str(header_row))

    def _write_ranking_block(
        self,
        ws: Any,
        mod: AnalyticsModule,
        *,
        chart: bool = True,
    ) -> None:
        items = list(mod.ranking)
        if not items:
            ws.append(["Sin ranking."])
            return

        ws.append(["Ranking"])
        ws.cell(row=ws.max_row, column=1).font = _SECTION_FONT
        header_row = ws.max_row + 1
        ws.append(["Posición", "Nombre", "Cantidad", "Participación %"])
        _style_header(ws, header_row)
        data_start = header_row + 1
        colors: list[str] = []
        for idx, item in enumerate(items):
            ws.append(
                [
                    item.rank,
                    item.label,
                    item.raw_value,
                    item.share_percentage,
                ]
            )
            if mod.id == "occupancy":
                colors.append(occupancy_color(item.share_percentage))
            else:
                colors.append(ranking_color(mod.id, idx))
        data_end = ws.max_row
        _try_table(ws, f"Ranking_{mod.id}", f"A{header_row}:D{data_end}")

        if not chart:
            return

        bar = BarChart()
        bar.type = "bar"
        bar.style = 10
        bar.title = mod.title
        bar.y_axis.title = None
        data_ref = Reference(ws, min_col=3, min_row=header_row, max_row=data_end)
        cats = Reference(ws, min_col=2, min_row=data_start, max_row=data_end)
        bar.add_data(data_ref, titles_from_data=True)
        bar.set_categories(cats)
        bar.shape = 4
        bar.height = 10
        bar.width = 16
        _apply_bar_colors(bar, colors)
        ws.add_chart(bar, "F" + str(header_row))

    def _write_hidden_meta(
        self,
        wb: Workbook,
        dashboard: DashboardResponse,
        *,
        generated_by: str,
    ) -> None:
        ws = wb.create_sheet("_sistema")
        ws.sheet_state = "hidden"
        ws.append(["clave", "valor"])
        ws.append(["schema_version", ANALYTICS_SCHEMA_VERSION])
        ws.append(["report_version", 2])
        ws.append(["generated_at", dashboard.generated_at.isoformat()])
        ws.append(["timezone", APP_TIMEZONE_NAME])
        ws.append(["generated_by", generated_by])
        ws.append(
            [
                "filters",
                (
                    f"range={dashboard.period.preset};"
                    f"from={dashboard.period.start};to={dashboard.period.end}"
                ),
            ]
        )
        for mod in dashboard.modules:
            ws.append([f"module:{mod.id}", mod.status.value])

    @staticmethod
    def visible_sheet_names(data: bytes) -> list[str]:
        wb = load_workbook(BytesIO(data))
        return [name for name in wb.sheetnames if not name.startswith("_")]

    @staticmethod
    def assert_no_forbidden_terms_in_visible(data: bytes) -> list[str]:
        """Return list of forbidden hits found in visible sheets (empty = ok)."""
        wb = load_workbook(BytesIO(data), data_only=False)
        hits: list[str] = []
        for name in wb.sheetnames:
            ws = wb[name]
            if getattr(ws, "sheet_state", "visible") == "hidden" or name.startswith("_"):
                continue
            for row in ws.iter_rows(values_only=True):
                for cell in row:
                    if cell is None:
                        continue
                    text = str(cell)
                    for term in FORBIDDEN_VISIBLE_TERMS:
                        if term in text:
                            hits.append(f"{name}:{term}")
        return hits

    @staticmethod
    def chart_count(data: bytes) -> int:
        wb = load_workbook(BytesIO(data))
        total = 0
        for name in wb.sheetnames:
            ws = wb[name]
            total += len(getattr(ws, "_charts", []) or [])
        return total
