"""Unit tests for analytics v2: formulas, countries, prefs, catalog, export."""

from __future__ import annotations

from datetime import date, datetime, timezone
from types import SimpleNamespace

import pytest

from app.common.enums import Permission, UserRole
from app.schemas.analytics_v2 import (
    ANALYTICS_SCHEMA_VERSION,
    AnalyticsModule,
    AnalyticsPreferencesUpdateSchema,
    BreakdownItem,
    ComparisonMode,
    DashboardQuery,
    DashboardResponse,
    DateRangePreset,
    Freshness,
    ModuleCategory,
    ModuleStatus,
    Period,
    PrimaryValue,
    RankingItem,
    Series,
    SeriesPoint,
    ValueType,
    VisualizationType,
    build_comparison,
    format_currency_cop,
    freshness_label,
)
from app.services.analytics_catalog_service import (
    ADMIN_DEFAULT_MODULES,
    GUIDE_DEFAULT_MODULES,
    LEGACY_LEAD_TO_MODULE,
    AnalyticsCatalogService,
)
from app.services.analytics_country_normalizer import AnalyticsCountryNormalizer
from app.services.analytics_export_service import AnalyticsExportService
from app.services.analytics_preferences_service import AnalyticsPreferencesService
from app.services.analytics_query_service import (
    AnalyticsQueryService,
    aggregation_grain,
    resolve_period,
)
from app.services.analytics_service import HOME_INELIGIBLE_IDS, _lead


# ── formulas / comparison ────────────────────────────────────────────────


def test_quoted_amount_never_labeled_facturado_in_helper() -> None:
    # format helper is currency; label is caller's responsibility — guard wording
    assert "facturado" not in format_currency_cop(1500000).lower()
    assert format_currency_cop(1500).startswith("$")


def test_build_comparison_uses_absolute_when_base_small() -> None:
    cmp = build_comparison(current=2, previous=1, previous_label="el periodo anterior")
    assert cmp.mode == ComparisonMode.ABSOLUTE
    assert cmp.sufficient_sample is False
    assert "Pasó de" in (cmp.label or "")


def test_build_comparison_percentage_when_base_sufficient() -> None:
    cmp = build_comparison(current=118, previous=100, previous_label="los 30 días anteriores")
    assert cmp.mode == ComparisonMode.BOTH
    assert cmp.percentage_delta == 18.0
    assert "18.0 % más" in (cmp.label or "") or "18 % más" in (cmp.label or "")


def test_conversion_blocked_documented_in_catalog_concepts() -> None:
    from app.services.analytics_catalog_service import BLOCKED_CONCEPTS

    assert "commercial_conversion" in BLOCKED_CONCEPTS
    assert "collected_revenue" in BLOCKED_CONCEPTS


def test_resolve_period_last_30_and_previous_equal_length() -> None:
    today = date(2026, 7, 14)
    current, previous = resolve_period(
        DashboardQuery(range=DateRangePreset.LAST_30_DAYS), today=today
    )
    assert current.end == today
    assert current.start == date(2026, 6, 15)
    assert (current.end - current.start).days == (previous.end - previous.start).days


def test_aggregation_grain() -> None:
    p7 = Period(start=date(2026, 7, 8), end=date(2026, 7, 14), label="7d")
    p60 = Period(start=date(2026, 5, 16), end=date(2026, 7, 14), label="2m")
    p365 = Period(start=date(2026, 1, 1), end=date(2026, 7, 14), label="y")
    assert aggregation_grain(p7) == "day"
    assert aggregation_grain(p60) == "week"
    assert aggregation_grain(p365) == "month"


def test_service_date_bounds_are_bson_encodable() -> None:
    """Raw Motor aggregates cannot encode datetime.date — use naive datetime."""
    from bson import encode

    start, end = AnalyticsQueryService._service_date_bounds(
        date(2026, 7, 14), date(2026, 8, 13)
    )
    assert isinstance(start, datetime)
    assert isinstance(end, datetime)
    encode({"requested_date": {"$gte": start, "$lte": end}})
    assert AnalyticsQueryService._format_service_date(start) == "2026-07-14"
    assert AnalyticsQueryService._format_service_date(date(2026, 8, 13)) == "2026-08-13"


def test_experience_capacity_prefers_standard_then_tiers() -> None:
    from types import SimpleNamespace

    assert AnalyticsQueryService._experience_capacity(None) is None
    assert (
        AnalyticsQueryService._experience_capacity(
            SimpleNamespace(standard_max_participants=10, base_capacity=8, pricing=None)
        )
        == 10
    )
    assert (
        AnalyticsQueryService._experience_capacity(
            SimpleNamespace(standard_max_participants=None, base_capacity=8, pricing=None)
        )
        == 8
    )
    tiers = [
        SimpleNamespace(max_participants=4),
        SimpleNamespace(max_participants=12),
    ]
    assert (
        AnalyticsQueryService._experience_capacity(
            SimpleNamespace(
                standard_max_participants=None,
                base_capacity=None,
                pricing=SimpleNamespace(tiers=tiers),
            )
        )
        == 12
    )


def test_equine_workload_reads_assignments_not_stale_field() -> None:
    """Workload must be computed from assignments, not the stale equine field."""
    import inspect

    src = inspect.getsource(AnalyticsQueryService._module_equine_workload)
    assert "workload_last_7_days" not in src
    assert "AssignmentDocument" in src
    assert "requested_date" in src
    assert "Collections.RESERVATIONS" in src


# ── countries ────────────────────────────────────────────────────────────


def test_country_normalizer_common_variants() -> None:
    n = AnalyticsCountryNormalizer()
    assert n.normalize("Colombia").country_code == "CO"
    assert n.normalize("USA").country_code == "US"
    assert n.normalize("Estados Unidos").country_code == "US"
    assert n.normalize("EEUU").country_code == "US"
    assert n.normalize("México").country_code == "MX"
    assert n.normalize("us").country_code == "US"
    assert n.normalize("Somalia").country_code == "SO"
    assert n.normalize("SO").country_code == "SO"
    assert n.normalize("so").country_name == "Somalia"


def test_country_normalizer_does_not_invent() -> None:
    n = AnalyticsCountryNormalizer()
    result = n.normalize("Narnia")
    assert result.resolved is False
    assert result.country_code is None
    assert result.original == "Narnia"


def test_country_normalizer_report_unresolved() -> None:
    n = AnalyticsCountryNormalizer()
    _, report = n.normalize_many(["Colombia", "Atlantis", "USA", "Atlantis"])
    assert report.resolved == 2
    assert report.unresolved == 2
    assert "Atlantis" in report.unresolved_values


# ── catalog / permissions ────────────────────────────────────────────────


def test_catalog_filters_payment_for_guide() -> None:
    catalog = AnalyticsCatalogService()
    admin = {m.id for m in catalog.get_catalog(UserRole.ADMIN).modules}
    guide = {m.id for m in catalog.get_catalog(UserRole.GUIDE).modules}
    assert "payment_status" in admin
    assert "payment_status" not in guide
    assert "reservation_trend" in guide
    assert "equine_care_alerts" in guide


def test_defaults_by_role() -> None:
    catalog = AnalyticsCatalogService()
    assert catalog.defaults_for_role(UserRole.ADMIN) == [
        m for m in ADMIN_DEFAULT_MODULES if m in catalog.allowed_module_ids(UserRole.ADMIN)
    ][:4]
    guide_defaults = catalog.defaults_for_role(UserRole.GUIDE)
    assert len(guide_defaults) <= 4
    for mid in guide_defaults:
        assert mid in GUIDE_DEFAULT_MODULES or mid in catalog.allowed_module_ids(UserRole.GUIDE)


def test_legacy_pin_mapping_skips_unknown() -> None:
    catalog = AnalyticsCatalogService()
    mapped = catalog.map_legacy_pins(
        ["ing_confirmed", "vol_conversion", "ori_top_pais", "unknown_x"],
        UserRole.ADMIN,
    )
    assert "confirmed_value_trend" in mapped
    assert "top_countries" in mapped
    assert "vol_conversion" not in LEGACY_LEAD_TO_MODULE or "vol_conversion" not in mapped


# ── preferences ──────────────────────────────────────────────────────────


def test_preferences_cap_at_four_and_drop_action_center() -> None:
    svc = AnalyticsPreferencesService()

    class FakeUser:
        role = UserRole.ADMIN
        analytics_preferences: dict = {}
        leads_preferences: dict = {}

        async def save(self) -> None:
            return None

    user = FakeUser()
    # sync path via _clean
    cleaned = svc._clean(
        {
            "schema_version": 2,
            "selected_module_ids": [
                "action_center",
                "reservation_trend",
                "confirmed_value_trend",
                "top_experiences",
                "top_countries",
                "occupancy",
            ],
        },
        UserRole.ADMIN,
    )
    assert "action_center" not in cleaned.selected_module_ids
    assert len(cleaned.selected_module_ids) <= 4


@pytest.mark.asyncio
async def test_preferences_update_persists() -> None:
    svc = AnalyticsPreferencesService()

    class FakeUser:
        role = UserRole.ADMIN
        analytics_preferences: dict = {}
        leads_preferences: dict = {}
        saved = False

        async def save(self) -> None:
            self.saved = True

    user = FakeUser()
    body = AnalyticsPreferencesUpdateSchema(
        selected_module_ids=[
            "confirmed_value_trend",
            "reservation_trend",
            "top_experiences",
            "top_countries",
            "occupancy",
        ]
    )
    result = await svc.update_preferences(user, body)  # type: ignore[arg-type]
    assert len(result.selected_module_ids) == 4
    assert user.saved is True
    assert user.analytics_preferences["schema_version"] == ANALYTICS_SCHEMA_VERSION


def test_lazy_migrate_from_legacy_pins() -> None:
    svc = AnalyticsPreferencesService()

    class FakeUser:
        role = UserRole.ADMIN
        analytics_preferences: dict = {}
        leads_preferences = {"pinned_lead_ids": ["ing_confirmed", "exp_top"]}

    prefs = svc.get_preferences(FakeUser())  # type: ignore[arg-type]
    assert prefs.selected_module_ids[0] == "confirmed_value_trend"
    assert "top_experiences" in prefs.selected_module_ids


# ── legacy helper still works for tests ──────────────────────────────────


def test_lead_helper_sets_home_eligible_from_blacklist() -> None:
    blocked = _lead(
        id="eq_total",
        category="catalogo",
        title="Total equinos",
        value="10",
        unit="equinos",
        description="x",
        icon="pets",
        order=1,
    )
    assert blocked.home_eligible is False
    assert "eq_total" in HOME_INELIGIBLE_IDS


# ── export ───────────────────────────────────────────────────────────────


def _sample_dashboard() -> DashboardResponse:
    now = datetime.now(timezone.utc)
    period = Period(
        start=date(2026, 6, 15),
        end=date(2026, 7, 14),
        label="Últimos 30 días",
        preset=DateRangePreset.LAST_30_DAYS,
    )
    fresh = Freshness(generated_at=now, label="Actualizado ahora")
    trend = AnalyticsModule(
        id="reservation_trend",
        category=ModuleCategory.RESERVATIONS,
        title="Tendencia de reservas",
        description="Reservas nuevas",
        visualization=VisualizationType.LINE,
        period=period,
        primary_value=PrimaryValue(
            raw=18, formatted="18", unit="reservas", value_type=ValueType.COUNT
        ),
        series=[
            Series(
                id="reservations",
                label="Reservas nuevas",
                unit="reservas",
                points=[
                    SeriesPoint(raw=3, label="2026-07-01", unit="reservas"),
                    SeriesPoint(raw=5, label="2026-07-02", unit="reservas"),
                ],
            )
        ],
        status=ModuleStatus.OK,
        insight_text="Las reservas aumentaron frente al periodo anterior.",
        generated_at=now,
        freshness=fresh,
    )
    origins = AnalyticsModule(
        id="reservation_origins",
        category=ModuleCategory.RESERVATIONS,
        title="Orígenes de reserva",
        description="Canales de origen",
        visualization=VisualizationType.DONUT,
        period=period,
        primary_value=PrimaryValue(
            raw=10, formatted="10", unit="reservas", value_type=ValueType.COUNT
        ),
        breakdown=[
            BreakdownItem(
                dimension="channel",
                key="whatsapp",
                label="WhatsApp",
                raw_value=6,
                formatted_value="6",
                unit="reservas",
                share_percentage=60.0,
            ),
            BreakdownItem(
                dimension="channel",
                key="instagram",
                label="Instagram",
                raw_value=4,
                formatted_value="4",
                unit="reservas",
                share_percentage=40.0,
            ),
        ],
        status=ModuleStatus.OK,
        insight_text="La mayoría de las reservas llegan por WhatsApp (60%, 6 reservas).",
        generated_at=now,
        freshness=fresh,
    )
    return DashboardResponse(
        modules=[trend, origins],
        period=period,
        generated_at=now,
        freshness=fresh,
        schema_version=ANALYTICS_SCHEMA_VERSION,
    )


def test_export_has_resumen_and_one_sheet_per_indicator() -> None:
    svc = AnalyticsExportService()
    data, filename = svc.export_dashboard(
        _sample_dashboard(), generated_by="Ana Admin", include_logo=False
    )
    assert filename.endswith(".xlsx")
    assert "analitica_" in filename
    names = AnalyticsExportService.visible_sheet_names(data)
    assert names[0] == "Resumen"
    assert "Tendencia de reservas" in names
    assert "Orígenes de reserva" in names
    assert "Indicadores" not in names
    assert "Tendencias" not in names
    assert "Desgloses" not in names
    hits = AnalyticsExportService.assert_no_forbidden_terms_in_visible(data)
    assert hits == [], hits
    assert AnalyticsExportService.chart_count(data) >= 2


def test_export_single_module_filter_shape() -> None:
    """When only one module is present (module_ids filter), Resumen + 1 sheet."""
    full = _sample_dashboard()
    single = DashboardResponse(
        modules=[full.modules[1]],  # origins only
        period=full.period,
        generated_at=full.generated_at,
        freshness=full.freshness,
        schema_version=full.schema_version,
    )
    svc = AnalyticsExportService()
    data, _ = svc.export_dashboard(single, generated_by="Ana", include_logo=False)
    names = AnalyticsExportService.visible_sheet_names(data)
    assert names == ["Resumen", "Orígenes de reserva"]
    assert AnalyticsExportService.chart_count(data) >= 1


def test_catalog_includes_reservation_origins() -> None:
    catalog = AnalyticsCatalogService().get_catalog(UserRole.ADMIN)
    ids = {m.id for m in catalog.modules}
    assert "reservation_origins" in ids
    origins = next(m for m in catalog.modules if m.id == "reservation_origins")
    assert origins.recommended_visualization == VisualizationType.DONUT
    assert origins.title == "Orígenes de reserva"


def test_channel_colors_canonical() -> None:
    from app.services.analytics_colors import CHANNEL_COLORS, breakdown_color

    assert CHANNEL_COLORS["whatsapp"] == "25D366"
    assert CHANNEL_COLORS["facebook"] == "1877F2"
    assert CHANNEL_COLORS["instagram"] == "E1306C"
    assert CHANNEL_COLORS["email"] == "64748B"
    assert breakdown_color("reservation_origins", "whatsapp", 0) == "25D366"


def test_rich_analysis_for_donut_and_line() -> None:
    from app.services.analytics_analysis import (
        analysis_parameter_rows,
        build_rich_analysis,
    )

    dash = _sample_dashboard()
    trend_lines = build_rich_analysis(dash.modules[0])
    assert any("promedio" in line.lower() for line in trend_lines)
    assert any("Rango:" in line or "mínimo" in line.lower() for line in trend_lines)
    assert any("Dispersión" in line or "Inicio→fin" in line or "Tendencia" in line for line in trend_lines)

    origins_lines = build_rich_analysis(dash.modules[1])
    assert any("WhatsApp" in line for line in origins_lines)
    assert any("HHI" in line or "Concentración" in line or "concentr" in line.lower() for line in origins_lines)
    assert any("Redes sociales" in line for line in origins_lines)

    trend_params = dict(analysis_parameter_rows(dash.modules[0]))
    assert "Puntos" in trend_params
    assert "Promedio" in trend_params
    origins_params = dict(analysis_parameter_rows(dash.modules[1]))
    assert "HHI" in origins_params
    assert "Líder" in origins_params


def test_export_includes_parameters_section() -> None:
    svc = AnalyticsExportService()
    data, _ = svc.export_dashboard(
        _sample_dashboard(), generated_by="Ana", include_logo=False
    )
    from io import BytesIO

    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(data))
    ws = wb["Orígenes de reserva"]
    values = [str(c.value) if c.value is not None else "" for row in ws.iter_rows() for c in row]
    assert any(v == "Parámetros" for v in values)
    assert any(v == "HHI" for v in values)
    assert any(v == "Análisis" for v in values)


def test_export_continues_without_logo() -> None:
    svc = AnalyticsExportService()
    data, _ = svc.export_dashboard(
        _sample_dashboard(), generated_by="Guía", include_logo=True
    )
    assert len(data) > 1000
    assert AnalyticsExportService.visible_sheet_names(data)[0] == "Resumen"


def test_freshness_label_business_language() -> None:
    now = datetime(2026, 7, 14, 12, 0, tzinfo=timezone.utc)
    assert freshness_label(now, now=now) == "Actualizado ahora"
    earlier = datetime(2026, 7, 14, 11, 52, tzinfo=timezone.utc)
    assert "minuto" in freshness_label(earlier, now=now)


def test_leads_preferences_caps_pins_and_drops_dupes() -> None:
    from app.schemas.analytics_v2 import LeadsPreferencesSchema

    schema = LeadsPreferencesSchema.from_user_prefs(
        {
            "pinned_lead_ids": ["a", "b", "a", "c", "d", "e", "f", "", 123],
            "excluded_lead_ids": ["x", "a", "y", "x"],
        }
    )
    assert schema.pinned_lead_ids == ["a", "b", "c", "d", "e"]
    assert schema.excluded_lead_ids == ["x", "y"]
